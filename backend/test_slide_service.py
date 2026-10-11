import os
import tempfile
import pytest
import numpy as np
import cv2
from pathlib import Path
from pptx import Presentation
from fastapi.testclient import TestClient

from app.main import app
from app.models.schemas import ChapterItem, ChapterSlideItem, ExtractSlidesRequest, VideoSlidesResponse
from app.services.slide_service import SlideService, slide_service


@pytest.fixture
def temp_slides_service():
    """Provides a SlideService instance using a temporary directory."""
    with tempfile.TemporaryDirectory() as td:
        service = SlideService(data_dir=Path(td))
        yield service


@pytest.fixture
def synthetic_video():
    """Creates a small 10-second synthetic test video file using OpenCV."""
    with tempfile.TemporaryDirectory() as td:
        video_path = os.path.join(td, "test_sample.mp4")
        fourcc = cv2.VideoWriter_fourcc(*'mp4v')
        out = cv2.VideoWriter(video_path, fourcc, 10.0, (320, 240))
        for sec in range(10):
            for f in range(10):
                img = np.zeros((240, 320, 3), dtype=np.uint8)
                cv2.putText(img, f"T={sec}.{f}s", (20, 100), cv2.FONT_HERSHEY_SIMPLEX, 0.8, (0, 255, 0), 2)
                out.write(img)
        out.release()
        yield video_path


class TestExtractionTimestampCalculation:
    def test_standard_chapter_calculation(self):
        """
        Chapter: 02:10 (130s) to 03:46 (226s)
        Expected: 226 - 5 = 221s (03:41)
        """
        ts = SlideService.compute_extraction_timestamp(start_time=130.0, end_time=226.0)
        assert ts == 221.0

    def test_short_chapter_under_five_seconds(self):
        """
        Chapter: 10.0s to 13.0s (duration 3.0s)
        Formula: max(10.0, 13.0 - 5.0) = max(10.0, 8.0) = 10.0s
        Extracts within its own time range instead of previous chapter.
        """
        ts = SlideService.compute_extraction_timestamp(start_time=10.0, end_time=13.0)
        assert ts == 10.0
        assert ts >= 10.0
        assert ts < 13.0

    def test_exact_five_second_chapter(self):
        """
        Chapter: 20.0s to 25.0s
        Formula: max(20.0, 25.0 - 5.0) = 20.0s
        """
        ts = SlideService.compute_extraction_timestamp(start_time=20.0, end_time=25.0)
        assert ts == 20.0

    def test_missing_or_invalid_end_time(self):
        """
        If end_time is None or <= start_time, safely defaults within chapter.
        """
        ts = SlideService.compute_extraction_timestamp(start_time=60.0, end_time=None)
        assert ts >= 60.0

    def test_video_duration_clamping(self):
        """
        Ensures timestamp does not exceed the total video duration.
        """
        ts = SlideService.compute_extraction_timestamp(start_time=90.0, end_time=100.0, total_duration=98.0)
        assert ts <= 97.5


class TestFrameExtractionAndPresentation:
    def test_extract_frame_from_synthetic_video(self, temp_slides_service, synthetic_video):
        """Tests that OpenCV successfully seeks and decodes a real video frame."""
        frame = temp_slides_service.extract_frame_from_video(synthetic_video, 4.0)
        assert frame is not None
        assert isinstance(frame, np.ndarray)
        assert frame.shape == (240, 320, 3)

    def test_extract_frame_nonexistent_video(self, temp_slides_service):
        """Gracefully returns None for missing video file."""
        frame = temp_slides_service.extract_frame_from_video("non_existent_file.mp4", 2.0)
        assert frame is None

    def test_pptx_generation_widescreen_and_validity(self, temp_slides_service):
        """
        Generates PPTX with multiple chapters and validates:
        1. 16:9 widescreen layout dimensions (13.333 x 7.5 inches).
        2. Exactly 1 slide per successfully extracted chapter.
        3. Correct ordering and title content.
        4. Valid, readable PPTX file.
        """
        with tempfile.TemporaryDirectory() as td:
            img_dir = Path(td)
            # Create 2 sample images
            for ch_id in ["ch_1", "ch_2"]:
                img = np.zeros((360, 640, 3), dtype=np.uint8)
                cv2.imwrite(str(img_dir / f"{ch_id}.jpg"), img)

            slides_data = [
                ChapterSlideItem(
                    chapter_id="ch_1",
                    chapter_index=1,
                    title="Introduction to Machine Learning",
                    start_time=0.0,
                    end_time=60.0,
                    duration=60.0,
                    timestamp_str="00:00 – 01:00",
                    extraction_timestamp=55.0,
                    extraction_timestamp_str="00:55",
                    image_url="/api/slides/test_vid/images/ch_1",
                    status="completed"
                ),
                ChapterSlideItem(
                    chapter_id="ch_2",
                    chapter_index=2,
                    title="Supervised vs Unsupervised Learning",
                    start_time=60.0,
                    end_time=150.0,
                    duration=90.0,
                    timestamp_str="01:00 – 02:30",
                    extraction_timestamp=145.0,
                    extraction_timestamp_str="02:25",
                    image_url="/api/slides/test_vid/images/ch_2",
                    status="completed"
                ),
                # A failed slide should not be included in the presentation
                ChapterSlideItem(
                    chapter_id="ch_3",
                    chapter_index=3,
                    title="Failed Chapter",
                    start_time=150.0,
                    end_time=200.0,
                    duration=50.0,
                    timestamp_str="02:30 – 03:20",
                    extraction_timestamp=195.0,
                    extraction_timestamp_str="03:15",
                    status="failed"
                )
            ]

            pptx_output = Path(td) / "test_out.pptx"
            success = temp_slides_service.generate_presentation(
                video_title="ML Course 101",
                slides_data=slides_data,
                images_dir=img_dir,
                output_pptx_path=pptx_output
            )
            assert success is True
            assert pptx_output.exists()
            assert pptx_output.stat().st_size > 5000

            # Validate by reading back with python-pptx
            prs = Presentation(str(pptx_output))
            assert len(prs.slides) == 2  # exactly 2 completed slides

            # Check 16:9 widescreen dimensions
            assert round(prs.slide_width.inches, 3) == 13.333
            assert round(prs.slide_height.inches, 3) == 7.5

            # Validate slide 1 content
            s1_texts = [shape.text for shape in prs.slides[0].shapes if shape.has_text_frame]
            assert any("Chapter 1: Introduction to Machine Learning" in t for t in s1_texts)
            assert any("00:55" in t for t in s1_texts)

            # Validate slide 2 content
            s2_texts = [shape.text for shape in prs.slides[1].shapes if shape.has_text_frame]
            assert any("Chapter 2: Supervised vs Unsupervised Learning" in t for t in s2_texts)
            assert any("02:25" in t for t in s2_texts)


class TestSlideOrchestrationAndEdgeCases:
    def test_missing_chapter_metadata_handling(self, temp_slides_service):
        """
        When video has no chapters, returns status 'no_chapters' with actionable message,
        without crashing or silently generating an incorrect PPT.
        """
        resp = temp_slides_service.extract_slides(
            video_id="no_ch_vid_123",
            chapters=[]
        )
        assert resp.status == "no_chapters"
        assert resp.has_chapters is False
        assert resp.total_chapters == 0
        assert resp.extracted_slides == 0
        assert "does not have chapter markers" in (resp.message or "")

    def test_full_pipeline_with_synthetic_video_override(self, temp_slides_service, synthetic_video):
        """
        Simulates end-to-end extraction using synthetic video frames without external YouTube download.
        """
        chapters = [
            ChapterItem(title="Intro", start_time=0.0, end_time=4.0, is_ai_detected=False),
            ChapterItem(title="Core", start_time=4.0, end_time=9.0, is_ai_detected=False)
        ]

        resp = temp_slides_service.extract_slides(
            video_id="demo_test_vid",
            chapters=chapters,
            video_path_override=synthetic_video
        )

        assert resp.status == "completed"
        assert resp.total_chapters == 2
        assert resp.extracted_slides == 2
        assert len(resp.slides) == 2
        assert resp.ppt_download_url == "/api/slides/demo_test_vid/download-ppt"

        # Check chapter 1 timestamp: end is 4.0, max(0.0, 4.0 - 5.0) = 0.0
        assert resp.slides[0].extraction_timestamp == 0.0
        # Check chapter 2 timestamp: end is 9.0, max(4.0, 9.0 - 5.0) = 4.0
        assert resp.slides[1].extraction_timestamp == 4.0

        # Verify PPT file exists
        ppt_path = temp_slides_service.get_presentation_path("demo_test_vid")
        assert ppt_path is not None and ppt_path.exists()

        # Verify cached retrieval
        cached = temp_slides_service.get_slides("demo_test_vid")
        assert cached is not None
        assert cached.extracted_slides == 2

    def test_security_path_traversal_prevention(self, temp_slides_service):
        """Ensures malicious directory traversal attempts are rejected."""
        with pytest.raises(ValueError):
            temp_slides_service.sanitize_id("../../../etc/passwd")

        with pytest.raises(ValueError):
            temp_slides_service.sanitize_id("vid; rm -rf /")


class TestSlidesApiEndpoints:
    def test_get_not_started_slides(self):
        client = TestClient(app)
        res = client.get("/api/slides/unextracted_vid_999")
        assert res.status_code == 200
        data = res.json()
        assert data["status"] == "not_started"
        assert data["extracted_slides"] == 0

    def test_extract_missing_video_id(self):
        client = TestClient(app)
        res = client.post("/api/slides/extract", json={"video_id": ""})
        assert res.status_code == 400

    def test_extract_empty_chapters_graceful_response(self):
        client = TestClient(app)
        res = client.post("/api/slides/extract", json={
            "video_id": "test_empty_ch",
            "chapters": []
        })
        assert res.status_code == 200
        data = res.json()
        assert data["status"] == "no_chapters"
        assert data["has_chapters"] is False

    def test_image_not_found(self):
        client = TestClient(app)
        res = client.get("/api/slides/nonexistent_vid/images/ch_1")
        assert res.status_code == 404

    def test_ppt_not_found(self):
        client = TestClient(app)
        res = client.get("/api/slides/nonexistent_vid/download-ppt")
        assert res.status_code == 404
