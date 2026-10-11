import os
import re
import json
import logging
import tempfile
from pathlib import Path
from typing import List, Optional, Dict, Any

import cv2
import numpy as np
from PIL import Image
import yt_dlp
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from app.models.schemas import ChapterItem, ChapterSlideItem, VideoSlidesResponse
from app.services.youtube_service import YouTubeService, youtube_service

logger = logging.getLogger(__name__)

# Base directory for persisted slides and presentations
_BACKEND_DIR = Path(__file__).resolve().parent.parent.parent
SLIDES_DATA_DIR = _BACKEND_DIR / "data" / "slides"
SLIDES_DATA_DIR.mkdir(parents=True, exist_ok=True)


class SlideService:
    def __init__(self, data_dir: Path = SLIDES_DATA_DIR):
        self.data_dir = data_dir.resolve()
        self.data_dir.mkdir(parents=True, exist_ok=True)

    @staticmethod
    def sanitize_id(identifier: str) -> str:
        """
        Sanitizes identifiers to prevent path traversal or unsafe characters.
        Allows alphanumeric, hyphen, and underscore.
        """
        if not identifier:
            raise ValueError("Identifier cannot be empty.")
        clean = identifier.strip()
        if not re.match(r'^[a-zA-Z0-9_-]{1,64}$', clean):
            raise ValueError(f"Invalid identifier format: '{identifier}'")
        return clean

    @staticmethod
    def sanitize_filename(filename: str) -> str:
        """
        Sanitizes a string to make it safe for filesystem and Content-Disposition header.
        """
        clean = re.sub(r'[\\/*?:"<>|]', "", filename)
        clean = clean.replace(" ", "_").strip("._")
        return clean[:100] or "presentation"

    @staticmethod
    def compute_extraction_timestamp(
        start_time: float,
        end_time: Optional[float],
        total_duration: Optional[float] = None
    ) -> float:
        """
        Calculates extraction timestamp for a chapter:
        Formula: extraction_timestamp = max(chapter_start_time, chapter_end_time - 5)
        
        Boundary guarantees:
        - If chapter is shorter than 5 seconds, extracts within its own time range
          rather than selecting a timestamp from previous chapter.
        - Clamps within video duration when known.
        - Preserves seconds internally with clean precision.
        """
        start = max(0.0, float(start_time))
        
        if end_time is not None and float(end_time) > start:
            end = float(end_time)
        else:
            end = start + 5.0

        # Exact formula: max(chapter_start_time, chapter_end_time - 5)
        extracted = max(start, end - 5.0)

        # Boundary safeguard for video duration
        if total_duration and total_duration > 0:
            extracted = min(extracted, max(start, float(total_duration) - 0.5))

        return round(extracted, 2)

    def get_video_dir(self, video_id: str) -> Path:
        """Returns safe resolved directory for a video's slides."""
        clean_id = self.sanitize_id(video_id)
        video_dir = (self.data_dir / clean_id).resolve()
        if not video_dir.is_relative_to(self.data_dir):
            raise ValueError(f"Path traversal detected for video_id: {video_id}")
        return video_dir

    def extract_frame_from_video(
        self,
        video_path: str,
        timestamp_seconds: float
    ) -> Optional[np.ndarray]:
        """
        Seeks to the calculated timestamp in seconds and decodes the frame using OpenCV.
        """
        if not os.path.exists(video_path):
            logger.error(f"Video file not found at: {video_path}")
            return None

        cap = cv2.VideoCapture(video_path)
        if not cap.isOpened():
            logger.error(f"Failed to open video file: {video_path}")
            return None

        try:
            # Seek in milliseconds
            target_ms = max(0.0, timestamp_seconds * 1000.0)
            cap.set(cv2.CAP_PROP_POS_MSEC, target_ms)

            ret, frame = cap.read()
            if ret and frame is not None and frame.size > 0:
                return frame

            # Fallback seek: slightly offset backward in case timestamp landed on the final keyframe boundary
            fallback_ms = max(0.0, (timestamp_seconds - 0.25) * 1000.0)
            cap.set(cv2.CAP_PROP_POS_MSEC, fallback_ms)
            ret, frame = cap.read()
            if ret and frame is not None and frame.size > 0:
                return frame

            logger.warning(f"Could not read frame at {timestamp_seconds}s from {video_path}")
            return None

        except Exception as e:
            logger.exception(f"Exception during frame extraction at {timestamp_seconds}s: {e}")
            return None
        finally:
            cap.release()

    def download_video_for_extraction(
        self,
        url_or_id: str,
        output_dir: str
    ) -> Optional[str]:
        """
        Downloads a lightweight video stream (e.g. 360p or lowest mp4) via yt-dlp
        into output_dir specifically for frame extraction.
        Avoids downloading gigantic 1080p/4K files.
        """
        url = url_or_id if url_or_id.startswith("http") else f"https://www.youtube.com/watch?v={url_or_id}"
        out_template = os.path.join(output_dir, "stream_video.%(ext)s")

        ydl_opts = {
            # Pick fast, lightweight 360p or 480p mp4 video stream (avc1/H.264 for optimal OpenCV compatibility)
            'format': 'bestvideo[height<=480][vcodec^=avc1]/bestvideo[height<=480]/worstvideo[vcodec^=avc1]/worstvideo/worst',
            'outtmpl': out_template,
            'quiet': True,
            'no_warnings': True,
            'noplaylist': True,
            'overwrites': True,
            'socket_timeout': 30,
        }

        try:
            logger.info(f"Downloading stream for frame extraction: {url}")
            with yt_dlp.YoutubeDL(ydl_opts) as ydl:
                ydl.download([url])

            # Find downloaded file
            for file in os.listdir(output_dir):
                if file.startswith("stream_video."):
                    full_path = os.path.join(output_dir, file)
                    if os.path.getsize(full_path) > 1024:
                        logger.info(f"Downloaded video stream ({os.path.getsize(full_path)} bytes) to: {full_path}")
                        return full_path

            logger.error("Download succeeded but no output file was located.")
            return None

        except Exception as e:
            logger.error(f"yt-dlp download failed for {url}: {e}")
            return None

    def generate_presentation(
        self,
        video_title: str,
        slides_data: List[ChapterSlideItem],
        images_dir: Path,
        output_pptx_path: Path
    ) -> bool:
        """
        Generates a widescreen 16:9 PowerPoint presentation (.pptx) containing
        one slide per successfully extracted chapter slide.
        Preserves image aspect ratio and uses modern typography and layout.
        """
        prs = Presentation()
        # Widescreen 16:9 standard dimensions
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        blank_layout = prs.slide_layouts[6]
        valid_slides = [s for s in slides_data if s.status == "completed"]

        if not valid_slides:
            logger.warning("No completed slides to generate PPT.")
            return False

        for current_idx, slide_item in enumerate(valid_slides, start=1):
            slide_item.slide_index = current_idx
            slide = prs.slides.add_slide(blank_layout)

            # Header Area
            header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.2))
            tf = header_box.text_frame
            tf.word_wrap = True
            tf.margin_left = Inches(0)
            tf.margin_right = Inches(0)
            tf.margin_top = Inches(0)
            tf.margin_bottom = Inches(0)

            # Paragraph 0: Chapter Number & Title
            p0 = tf.paragraphs[0]
            p0.text = f"Chapter {slide_item.chapter_index}: {slide_item.title}"
            p0.font.size = Pt(20)
            p0.font.bold = True
            p0.font.color.rgb = RGBColor(15, 23, 42)  # Slate 900

            # Paragraph 1: Video Timestamp and Slide Extraction Timestamp
            p1 = tf.add_paragraph()
            p1.text = (
                f"Chapter Timestamp: {slide_item.timestamp_str}   •   "
                f"Slide Frame at: {slide_item.extraction_timestamp_str} (5s before chapter end)"
            )
            p1.font.size = Pt(11)
            p1.font.color.rgb = RGBColor(100, 116, 139)  # Slate 500
            p1.space_before = Pt(4)

            # Image Insertion with Aspect Ratio Preservation
            img_filename = f"{slide_item.chapter_id}.jpg"
            img_path = (images_dir / img_filename).resolve()

            if img_path.exists():
                try:
                    with Image.open(img_path) as pil_img:
                        img_w, img_h = pil_img.size

                    aspect = img_w / img_h
                    max_w = 11.733
                    max_h = 5.2

                    if aspect > (max_w / max_h):
                        target_w = max_w
                        target_h = target_w / aspect
                    else:
                        target_h = max_h
                        target_w = target_h * aspect

                    pos_left = 0.8 + (max_w - target_w) / 2.0
                    pos_top = 1.75 + (max_h - target_h) / 2.0

                    slide.shapes.add_picture(
                        str(img_path),
                        Inches(pos_left),
                        Inches(pos_top),
                        Inches(target_w),
                        Inches(target_h)
                    )
                except Exception as img_err:
                    logger.warning(f"Could not insert image {img_path} into slide: {img_err}")

            # Subtle Footer
            footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.35))
            ftf = footer_box.text_frame
            fp = ftf.paragraphs[0]
            fp.text = f"{video_title[:75]}  •  Generated by NoteTube"
            fp.font.size = Pt(9)
            fp.font.color.rgb = RGBColor(148, 163, 184)  # Slate 400

        output_pptx_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(output_pptx_path))
        logger.info(f"Successfully generated PowerPoint at: {output_pptx_path}")
        return True

    def extract_slides(
        self,
        video_id: str,
        youtube_url: Optional[str] = None,
        chapters: Optional[List[ChapterItem]] = None,
        force: bool = False,
        video_path_override: Optional[str] = None  # Useful for offline tests
    ) -> VideoSlidesResponse:
        """
        Orchestrates full chapter-to-slide extraction and PowerPoint presentation creation:
        1. Validates video identifier and ensures isolated durable storage.
        2. Retrieves or parses chapter metadata (gracefully detects if chapters exist).
        3. Computes extraction timestamp for each chapter (5 seconds before chapter end).
        4. Downloads/accesses video stream and extracts frames using OpenCV.
        5. Saves images deterministically associated with stable chapter IDs.
        6. Generates high-quality widescreen PPTX file.
        7. Saves metadata and returns structured VideoSlidesResponse.
        """
        clean_id = self.sanitize_id(video_id)
        video_dir = self.get_video_dir(clean_id)
        images_dir = video_dir / "images"
        meta_file = video_dir / "slides_meta.json"
        ppt_file = video_dir / "presentation.pptx"

        # Check existing cache unless force requested
        if not force and meta_file.exists() and ppt_file.exists():
            try:
                with open(meta_file, "r", encoding="utf-8") as f:
                    cached_data = json.load(f)
                if cached_data.get("status") in ["completed", "no_chapters"]:
                    logger.info(f"Returning cached slides metadata for video: {clean_id}")
                    return VideoSlidesResponse(**cached_data)
            except Exception as e:
                logger.warning(f"Failed to read cached slides metadata: {e}. Re-extracting...")

        video_dir.mkdir(parents=True, exist_ok=True)
        images_dir.mkdir(parents=True, exist_ok=True)

        # 1. Resolve video metadata and chapter information
        url = youtube_url or f"https://www.youtube.com/watch?v={clean_id}"
        resolved_chapters: List[ChapterItem] = []
        video_title = "YouTube Video"
        total_duration: Optional[float] = None

        if chapters and len(chapters) > 0:
            resolved_chapters = chapters
        else:
            try:
                meta = youtube_service.get_video_metadata(url)
                video_title = meta.title
                total_duration = float(meta.duration) if meta.duration else None
                if meta.chapters and len(meta.chapters) > 0:
                    resolved_chapters = meta.chapters
            except Exception as e:
                logger.warning(f"Could not retrieve metadata for {url}: {e}")

        # 2. Handle videos with no chapters gracefully
        if not resolved_chapters or len(resolved_chapters) == 0:
            logger.info(f"No chapter metadata available for video {clean_id}.")
            response = VideoSlidesResponse(
                video_id=clean_id,
                video_title=video_title,
                status="no_chapters",
                has_chapters=False,
                message="This video does not have chapter markers. Chapter slides and presentation export require video chapters.",
                total_chapters=0,
                extracted_slides=0,
                slides=[],
                ppt_download_url=None
            )
            with open(meta_file, "w", encoding="utf-8") as f:
                json.dump(response.model_dump(), f, indent=2)
            return response

        # 3. Compute extraction timestamps for all chapters
        slide_items: List[ChapterSlideItem] = []
        for idx, ch in enumerate(resolved_chapters):
            ch_num = idx + 1
            ch_id = ch.chapter_id or f"ch_{ch_num}"
            start_t = max(0.0, ch.start_time)
            
            if ch.end_time is not None and ch.end_time > start_t:
                end_t = ch.end_time
            elif ch.duration is not None and ch.duration > 0:
                end_t = start_t + ch.duration
            elif idx + 1 < len(resolved_chapters):
                end_t = resolved_chapters[idx + 1].start_time
            elif total_duration and total_duration > start_t:
                end_t = total_duration
            else:
                end_t = start_t + 60.0

            dur = max(0.0, end_t - start_t)
            extraction_t = self.compute_extraction_timestamp(start_t, end_t, total_duration)

            start_str = YouTubeService.format_timestamp(start_t)
            end_str = YouTubeService.format_timestamp(end_t)
            t_str = ch.timestamp_str or f"{start_str} – {end_str}"
            ext_t_str = YouTubeService.format_timestamp(extraction_t)

            slide_items.append(ChapterSlideItem(
                chapter_id=ch_id,
                chapter_index=ch_num,
                title=ch.title,
                start_time=start_t,
                end_time=end_t,
                duration=dur,
                timestamp_str=t_str,
                extraction_timestamp=extraction_t,
                extraction_timestamp_str=ext_t_str,
                image_url=f"/api/slides/{clean_id}/images/{ch_id}",
                status="pending"
            ))

        # 4. Extract frames
        extracted_count = 0

        # Video source: either an override path (for test mocking) or downloaded stream
        if video_path_override and os.path.exists(video_path_override):
            for item in slide_items:
                frame = self.extract_frame_from_video(video_path_override, item.extraction_timestamp)
                if frame is not None:
                    out_img = images_dir / f"{item.chapter_id}.jpg"
                    cv2.imwrite(str(out_img), frame)
                    item.status = "completed"
                    extracted_count += 1
                else:
                    item.status = "failed"
                    item.error = f"Frame could not be decoded at timestamp {item.extraction_timestamp_str}"
        else:
            with tempfile.TemporaryDirectory() as temp_dir:
                video_file = self.download_video_for_extraction(url, temp_dir)
                if not video_file or not os.path.exists(video_file):
                    logger.error(f"Video download failed for {clean_id}. Cannot extract frames.")
                    for item in slide_items:
                        item.status = "failed"
                        item.error = "Video download failed or video unavailable."

                    response = VideoSlidesResponse(
                        video_id=clean_id,
                        video_title=video_title,
                        status="failed",
                        has_chapters=True,
                        message="Failed to download video stream for frame extraction.",
                        total_chapters=len(slide_items),
                        extracted_slides=0,
                        slides=slide_items,
                        error="Video stream could not be downloaded."
                    )
                    with open(meta_file, "w", encoding="utf-8") as f:
                        json.dump(response.model_dump(), f, indent=2)
                    return response

                # Extract each chapter frame
                for item in slide_items:
                    try:
                        frame = self.extract_frame_from_video(video_file, item.extraction_timestamp)
                        if frame is not None and frame.size > 0:
                            out_img = images_dir / f"{item.chapter_id}.jpg"
                            cv2.imwrite(str(out_img), frame)
                            item.status = "completed"
                            extracted_count += 1
                        else:
                            item.status = "failed"
                            item.error = f"Could not decode frame at {item.extraction_timestamp_str}"
                    except Exception as ex:
                        logger.warning(f"Error extracting slide for {item.chapter_id}: {ex}")
                        item.status = "failed"
                        item.error = str(ex)

        # 5. Generate PowerPoint presentation
        ppt_url = None
        if extracted_count > 0:
            success = self.generate_presentation(
                video_title=video_title,
                slides_data=slide_items,
                images_dir=images_dir,
                output_pptx_path=ppt_file
            )
            if success:
                ppt_url = f"/api/slides/{clean_id}/download-ppt"

        # 6. Save metadata
        status = "completed" if extracted_count > 0 else "failed"
        response = VideoSlidesResponse(
            video_id=clean_id,
            video_title=video_title,
            status=status,
            has_chapters=True,
            total_chapters=len(slide_items),
            extracted_slides=extracted_count,
            slides=slide_items,
            ppt_download_url=ppt_url,
            message=f"Successfully extracted {extracted_count} of {len(slide_items)} chapter slides." if extracted_count > 0 else "No slide frames could be decoded."
        )

        with open(meta_file, "w", encoding="utf-8") as f:
            json.dump(response.model_dump(), f, indent=2)

        return response

    def get_slides(self, video_id: str) -> Optional[VideoSlidesResponse]:
        """Retrieves cached slides metadata for video_id if available."""
        clean_id = self.sanitize_id(video_id)
        video_dir = self.get_video_dir(clean_id)
        meta_file = video_dir / "slides_meta.json"

        if meta_file.exists():
            try:
                with open(meta_file, "r", encoding="utf-8") as f:
                    data = json.load(f)
                return VideoSlidesResponse(**data)
            except Exception as e:
                logger.error(f"Error reading slides metadata for {clean_id}: {e}")
        return None

    def get_slide_image_path(self, video_id: str, chapter_id: str) -> Optional[Path]:
        """Safely returns absolute Path to chapter image file with fallback name resolution."""
        clean_vid = self.sanitize_id(video_id)
        clean_cid = self.sanitize_id(chapter_id)
        images_dir = self.get_video_dir(clean_vid) / "images"
        
        img_path = (images_dir / f"{clean_cid}.jpg").resolve()
        if img_path.exists() and img_path.is_relative_to(self.data_dir):
            return img_path

        if images_dir.exists():
            alt1 = (images_dir / f"ch_{clean_cid}.jpg").resolve()
            if alt1.exists() and alt1.is_relative_to(self.data_dir):
                return alt1
            if clean_cid.startswith("ch_"):
                alt2 = (images_dir / f"{clean_cid[3:]}.jpg").resolve()
                if alt2.exists() and alt2.is_relative_to(self.data_dir):
                    return alt2
        return None

    def get_presentation_path(self, video_id: str) -> Optional[Path]:
        """Safely returns absolute Path to PowerPoint file."""
        clean_vid = self.sanitize_id(video_id)
        ppt_path = (self.get_video_dir(clean_vid) / "presentation.pptx").resolve()

        if ppt_path.exists() and ppt_path.is_relative_to(self.data_dir):
            return ppt_path
        return None


slide_service = SlideService()
